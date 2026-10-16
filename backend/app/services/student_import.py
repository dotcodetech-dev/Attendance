"""Bulk student import from an Excel workbook or CSV file - the student twin of employee_import.

Two entry points used by the API:
- inspect(): find the header row (the sheet may start with title rows), list the columns with sample
  values, guess which column is which field, and suggest the organization the students belong to.
- process(): with the confirmed mapping and organization, check every row and - when commit=True -
  create / update students (by Face ID = StudentCode). Missing schools and departments are created in
  that organization, linked Organization > School > Department. process(commit=False) runs the same
  steps and rolls back: the preview shows exactly what Import will do.

A Face ID that already belongs to a staff record:
- the readers created it themselves (name = the ID, or no department): the student is imported; the
  staff duplicate is listed so it can be removed with Employees -> "Clean up student records";
- a real staff member (e.g. a part-time scholar who is also staff): the row is skipped, the staff
  record keeps the ID.
"""

from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any, Optional

from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.models.entities import DepartmentMaster, Employee, Organization, School, Student, StudentShift

MAX_HEADER_SCAN_ROWS = 25
SAMPLE_COUNT = 3

FIELDS = ["StudentCode", "StudentName", "RollNo", "Department", "School", "RegisteredYear", "Batch",
          "Category", "Section", "RegisteredDate", "Status"]

# field -> normalized header aliases. "Access Card No" is the Face ID on these sheets; "Employee No" the roll number.
_ALIASES: dict[str, list[str]] = {
    "StudentCode": ["accesscardno", "cardno", "faceid", "deviceid", "biometricid", "studentcode", "code", "pin", "userid"],
    "StudentName": ["studentname", "employeename", "scholarname", "name", "fullname"],
    "RollNo": ["rollno", "rollnumber", "employeeno", "registerno", "regno", "registrationno", "studentno",
               "admissionno", "enrollmentno", "enrolmentno"],
    "Department": ["department", "dept", "departmentname"],
    "School": ["school", "schoolname", "faculty"],
    "RegisteredYear": ["registeredyear", "registeredsession", "session", "admissionsession", "subdivision"],
    "Batch": ["batch", "admissionyear", "division"],
    "Category": ["category", "mode", "studycategory"],
    "Section": ["section"],
    "RegisteredDate": ["dateofreg", "dateofregistration", "registereddate", "regdate", "dateofadmission",
                       "admissiondate", "doj", "dateofjoining"],
    "Status": ["status", "currentstatus"],
}
_EMPTY = {"", "na", "n/a", "nil", "none", "-", "--", "null"}
_MONTH_NAMES = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]
_DATE_FORMATS = ("%d-%b-%Y", "%d-%b-%y", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y")


def _hnorm(text: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _blank(text: str) -> bool:
    return text.strip().lower() in _EMPTY


def _name_key(text: Any) -> str:
    """School / department names compared ignoring case, '&' vs 'and', punctuation and extra spaces."""
    t = str(text or "").lower().replace("&", " and ")
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _acronym(text: str) -> str:
    """'School of Physical Sciences & Computational Sciences' -> 'pscs' (matches a school named 'School of PSCS')."""
    words = [w for w in _name_key(text).split() if w not in {"school", "of", "and", "the", "for"}]
    return "".join(w[0] for w in words)


def _school_core(text: str) -> str:
    return " ".join(w for w in _name_key(text).split() if w not in {"school", "of"})


def _reg_year(text: str) -> Optional[str]:
    """'July 2021 Registered' / 'Jul-2021' / 'jan 2026' -> 'Jul 2021' (the Students page's session format)."""
    t = text.strip()
    if not t:
        return None
    m = re.search(r"([A-Za-z]{3,9})[\s\-/,]*(\d{4})", t)
    if m and m.group(1)[:3].lower() in _MONTH_NAMES:
        return f"{m.group(1)[:3].title()} {m.group(2)}"
    return t[:20]


def _category(text: str) -> Optional[str]:
    t = text.strip()
    if not t:
        return None
    low = t.lower()
    if "full" in low:
        return "Full Time"
    if "part" in low:
        return "Part Time"
    return t[:20]


def _status(text: str) -> Optional[str]:
    low = text.strip().lower()
    if not low:
        return None
    if low.startswith(("discontinu", "inactive", "left", "drop", "withdraw", "cancel")):
        return "Discontinued"
    if low.startswith(("complet", "awarded", "passed", "graduat")):
        return "Completed"
    return "Active"


def _iso_date(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def read_rows(data: bytes, filename: str) -> list[list[Any]]:
    """All rows of the first sheet (.xlsx) or of a CSV file, as lists of cell values."""
    name = (filename or "").lower()
    if name.endswith(".csv") or not data.startswith(b"PK"):
        if name.endswith(".xls"):
            raise HTTPException(400, "Old .xls files are not supported: open it in Excel and save as .xlsx (or CSV).")
        for enc in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return [row for row in csv.reader(io.StringIO(text), dialect)]
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not read the file as Excel (.xlsx): {exc.__class__.__name__}")
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    wb.close()
    return rows


class StudentImportService:
    def __init__(self, db: Session) -> None:
        self.db = db

    # ---- inspect ----

    def _detect_header_row(self, rows: list[list[Any]]) -> int:
        aliases = {a for v in _ALIASES.values() for a in v}
        best, best_score = 0, -1
        for i, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
            score = sum(1 for c in row if _hnorm(c) in aliases)
            if score > best_score:
                best, best_score = i, score
        return best + 1

    def _samples(self, rows: list[list[Any]], header_row: int, idx: int, n: int = SAMPLE_COUNT) -> list[str]:
        out = []
        for r in rows[header_row:]:
            if idx < len(r):
                s = _cell(r[idx])
                if s and not _blank(s):
                    out.append(s)
            if len(out) >= n:
                break
        return out

    def guess_mapping(self, rows: list[list[Any]], header_row: int) -> dict[str, int]:
        header = rows[header_row - 1] if 0 < header_row <= len(rows) else []
        keys = [_hnorm(c) for c in header]
        mapping: dict[str, int] = {}
        for field, aliases in _ALIASES.items():   # alias order = preference (Access Card No before Code ...)
            for alias in aliases:
                idx = next((i for i, k in enumerate(keys) if k == alias and i not in mapping.values()), None)
                if idx is not None:
                    mapping[field] = idx
                    break
        # A "Section" column that holds school names ("School of ...") is the school.
        sec = mapping.get("Section")
        if sec is not None:
            samples = self._samples(rows, header_row, sec, 10)
            if samples and sum(1 for s in samples if s.lower().startswith("school")) >= len(samples) / 2:
                mapping.pop("Section")
                if "School" not in mapping:
                    mapping["School"] = sec
        # Prefer the registration date over a birth date for "Registered Date"; fall back to D.O.B.
        if "RegisteredDate" not in mapping:
            dob = next((i for i, k in enumerate(keys) if k in {"dob", "dateofbirth"} and i not in mapping.values()), None)
            if dob is not None:
                mapping["RegisteredDate"] = dob
        return mapping

    def suggest_org(self, rows: list[list[Any]], header_row: int, mapping: dict[str, int]) -> Optional[int]:
        orgs = self.db.query(Organization).filter((Organization.IsActive.is_(None)) | (Organization.IsActive == "Y")).all()
        if not orgs:
            return None
        body = rows[header_row:header_row + 200]
        text = " ".join(_cell(c).lower() for r in body for c in r)
        scholar_org = next((o for o in orgs if "scholar" in _name_key(o.OrgName)), None)
        if scholar_org and ("ph.d" in text or "phd" in text or "scholar" in text):
            return scholar_org.OrganizationId
        di = mapping.get("Department")
        if di is None:
            return orgs[0].OrganizationId
        names = {_name_key(_cell(r[di])) for r in body if di < len(r) and _cell(r[di])}
        best, best_n = orgs[0].OrganizationId, -1
        for o in orgs:
            n = sum(1 for d in self.db.query(DepartmentMaster).filter(DepartmentMaster.OrganizationId == o.OrganizationId)
                    if _name_key(d.DepartmentName) in names)
            if n > best_n:
                best, best_n = o.OrganizationId, n
        return best

    def inspect(self, data: bytes, filename: str) -> dict:
        rows = read_rows(data, filename)
        if not rows:
            raise HTTPException(400, "The file is empty.")
        header_row = self._detect_header_row(rows)
        header = rows[header_row - 1]
        ncols = max(len(r) for r in rows[:max(header_row + 50, 1)])
        columns = [{"index": i, "letter": get_column_letter(i + 1),
                    "header": (_cell(header[i]) if i < len(header) else "") or None,
                    "samples": self._samples(rows, header_row, i)} for i in range(ncols)]
        mapping = self.guess_mapping(rows, header_row)
        return {"header_row": header_row, "columns": columns, "guessed_mapping": mapping, "fields": FIELDS,
                "suggested_organization_id": self.suggest_org(rows, header_row, mapping)}

    # ---- masters ----

    def _load_masters(self, org_id: int) -> None:
        self.schools = self.db.query(School).all()
        self.depts = self.db.query(DepartmentMaster).filter(DepartmentMaster.OrganizationId == org_id).all()

    def _school(self, org_id: int, name: str, created: list[str]) -> Optional[School]:
        if not name or _blank(name):
            return None
        key, acr, core = _name_key(name), _acronym(name), _school_core(name)
        mine = [s for s in self.schools if s.OrganizationId == org_id]
        free = [s for s in self.schools if s.OrganizationId is None]
        for pool, link in ((mine, False), (free, True)):
            hit = (next((s for s in pool if _name_key(s.SchoolName) == key), None)
                   or next((s for s in pool if _school_core(s.SchoolName) in {acr, core} and len(acr) >= 2), None))
            if hit is not None:
                if link:
                    hit.OrganizationId = org_id
                return hit
        row = School(SchoolName=name.strip()[:150], IsActive="Y", OrganizationId=org_id)
        self.db.add(row)
        self.db.flush()
        self.schools.append(row)
        created.append(row.SchoolName)
        return row

    def _department(self, org_id: int, name: str, school: Optional[School], created: list[str]) -> DepartmentMaster:
        key = _name_key(name)
        same = [d for d in self.depts if _name_key(d.DepartmentName) == key]
        if school is not None:
            hit = next((d for d in same if d.SchoolId == school.SchoolId), None)
            if hit is None:
                hit = next((d for d in same if d.SchoolId is None), None)
                if hit is not None:
                    hit.SchoolId = school.SchoolId
            if hit is not None:
                return hit
        elif same:
            return same[0]
        row = DepartmentMaster(OrganizationId=org_id, DepartmentName=name.strip()[:150], IsActive="Y",
                               SchoolId=school.SchoolId if school is not None else None)
        self.db.add(row)
        self.db.flush()
        self.depts.append(row)
        created.append(row.DepartmentName + (f" ({school.SchoolName})" if school is not None else ""))
        return row

    # ---- process ----

    def process(self, data: bytes, filename: str, mapping: dict[str, int], header_row: int,
                organization_id: Optional[int], commit: bool) -> dict:
        org = self.db.get(Organization, organization_id) if organization_id else None
        if org is None:
            raise HTTPException(400, "Choose the organization these students belong to.")
        for f in ("StudentCode", "StudentName", "Department"):
            if f not in mapping:
                raise HTTPException(400, f'Map a column to "{f}" first.')
        rows = read_rows(data, filename)
        self._load_masters(org.OrganizationId)
        shift = (self.db.query(StudentShift.StudentShiftId).filter(StudentShift.IsActive == "Y")
                 .order_by(StudentShift.StudentShiftId).first())
        shift_id = shift[0] if shift else None
        students = {s.StudentCode: s for s in self.db.query(Student)}
        staff = {e.EmployeeCode: e for e in self.db.query(Employee)}
        created_schools: list[str] = []
        created_depts: list[str] = []
        seen: set[str] = set()
        out_rows: list[dict] = []
        summary = {"total": 0, "to_create": 0, "to_update": 0, "skipped": 0, "errors": 0, "staff_duplicates": 0}

        def get(r: list, field: str) -> str:
            i = mapping.get(field)
            return _cell(r[i]) if i is not None and i < len(r) else ""

        try:
            for n, r in enumerate(rows[header_row:], start=header_row + 1):
                if not any(_cell(c) for c in r):
                    continue
                code, name, dept_name = get(r, "StudentCode"), get(r, "StudentName"), get(r, "Department")
                info = {"row_number": n, "StudentCode": code or None, "StudentName": name or None,
                        "RollNo": get(r, "RollNo") or None, "Department": dept_name or None,
                        "School": (get(r, "School") if not _blank(get(r, "School")) else "") or None, "messages": []}
                summary["total"] += 1
                problems = []
                if not code or _blank(code):
                    problems.append("no Face ID (code)")
                elif len(code) > 50:
                    problems.append("Face ID longer than 50 characters")
                if not name:
                    problems.append("no name")
                if not dept_name or _blank(dept_name):
                    problems.append("no department")
                if code in seen:
                    problems.append("this Face ID appears twice in the file")
                if problems:
                    info.update(action="error", messages=problems)
                    summary["errors"] += 1
                    out_rows.append(info)
                    continue
                seen.add(code)
                emp = staff.get(code)
                student = students.get(code)
                if emp is not None and student is None:
                    auto = (emp.EmployeeName or "").strip() == code or not (emp.Department or "").strip()
                    if not auto:
                        info.update(action="skip", messages=[
                            f"Face ID belongs to staff member {emp.EmployeeName}"
                            + (f" ({emp.Department})" if emp.Department else "")
                            + " - kept as staff (e.g. a scholar who is also staff)"])
                        summary["skipped"] += 1
                        out_rows.append(info)
                        continue
                    info["messages"].append("also a reader-created staff record: remove it with Employees -> Clean up student records")
                    summary["staff_duplicates"] += 1
                school = self._school(org.OrganizationId, get(r, "School"), created_schools)
                dept = self._department(org.OrganizationId, dept_name, school, created_depts)
                values = {
                    "StudentName": name[:250],
                    "DepartmentId": dept.DepartmentId,
                    "SchoolId": school.SchoolId if school is not None else None,
                    "RollNo": (get(r, "RollNo") or None) and get(r, "RollNo")[:50],
                    "RegisteredYear": _reg_year(get(r, "RegisteredYear")),
                    "Batch": (get(r, "Batch") or None) and get(r, "Batch")[:20],
                    "Category": _category(get(r, "Category")),
                    "Section": (get(r, "Section") if not _blank(get(r, "Section")) else "")[:20] or None,
                    "DateOfBirth": _iso_date(r[mapping["RegisteredDate"]]) if "RegisteredDate" in mapping and mapping["RegisteredDate"] < len(r) else None,
                    "Status": _status(get(r, "Status")),
                }
                if student is None:
                    student = Student(StudentCode=code, YearOfStudy=1, StudentShiftId=shift_id, IsActive="Y",
                                      **{"Status": "Active", **{k: v for k, v in values.items() if v is not None}})
                    if student.Status != "Active":
                        student.IsActive = "N"
                        student.StatusDate = date.today().isoformat()
                    self.db.add(student)
                    students[code] = student
                    info["action"] = "create"
                    summary["to_create"] += 1
                else:
                    for k, v in values.items():
                        if v is not None and (k in ("StudentName", "DepartmentId") or f_mapped(k, mapping)):
                            setattr(student, k, v)
                    info["action"] = "update"
                    summary["to_update"] += 1
                out_rows.append(info)
            self.db.flush()
            if commit:
                self.db.commit()
            else:
                self.db.rollback()
        except Exception:
            self.db.rollback()
            raise
        if commit:
            summary["created"], summary["updated"] = summary["to_create"], summary["to_update"]
        return {"committed": commit, "organization": org.OrgName, "summary": summary, "rows": out_rows,
                "masters_created": {"schools": created_schools, "departments": created_depts}}


def f_mapped(field: str, mapping: dict[str, int]) -> bool:
    """On an update, only fields the sheet actually has are changed."""
    return {"DateOfBirth": "RegisteredDate"}.get(field, field) in mapping or (field == "SchoolId" and "School" in mapping)
