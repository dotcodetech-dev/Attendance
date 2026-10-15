"""
CSV, Excel and PDF rendering for HR reports. Takes the dicts produced by
app/services/reports.py and turns them into downloadable files.
Uses stdlib csv, openpyxl (xlsx) and reportlab (pdf).

Grouped reports (by Department) are expressed as ordinary data rows
interleaved with `Group` markers; every writer renders a marker as a spanning
section header so all three formats share one row builder.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any, Callable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


@dataclass
class Group:
    """Section header row. level 0 = department."""
    label: str
    level: int = 0


def _rows(data: dict, row_fn: Callable[[dict], list[Any]], grouped: bool) -> list[Any]:
    """Flatten a report into writer rows; with grouped=True insert Group markers
    from data['groups'] (one section per Department, with its count)."""
    if not grouped or not data.get("groups"):
        # no department sections: S.No runs 1..n over the whole list
        return [row_fn({**e, "SNo": i} if "SNo" in e else e) for i, e in enumerate(data["employees"], 1)]
    out: list[Any] = []
    for g in data["groups"]:
        out.append(Group(f"{g['Department']}  ({g['Count']})", 0))
        out += [row_fn(e) for e in g["employees"]]
    return out


# ---------------------------------------------------------------- CSV

def _csv(headers: list[str], rows: list[Any]) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(headers)
    for r in rows:
        w.writerow([("" if r.level else "") + r.label] if isinstance(r, Group) else r)
    return buf.getvalue()


# ---------------------------------------------------------------- Excel

_ILLEGAL_XL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _xl_value(v: Any) -> Any:
    """Excel refuses control characters; strip them (stray device data)."""
    return _ILLEGAL_XL.sub("", v) if isinstance(v, str) else v


_XL_DEPT = PatternFill("solid", fgColor="5B21B6")
_XL_DESIG = PatternFill("solid", fgColor="EDE9FE")
_XL_HEAD = PatternFill("solid", fgColor="4C1D95")
_XL_BODY = Font(size=12)  # report text was Excel's 11pt default; one step larger now


def _xlsx(title: str, subtitle: str, headers: list[str], rows: list[Any], org: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ncols = len(headers)
    r = 1
    if org and org.get("OrgName"):
        ws.cell(r, 1, org["OrgName"]).font = Font(bold=True, size=14, color="5B21B6")
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        r += 1
    ws.cell(r, 1, title).font = Font(bold=True, size=13)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    r += 1
    ws.cell(r, 1, subtitle).font = Font(italic=True, size=12, color="666666")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
    r += 2
    for c, h in enumerate(headers, 1):
        cell = ws.cell(r, c, h)
        cell.font = Font(bold=True, size=12, color="FFFFFF")
        cell.fill = _XL_HEAD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(r + 1, 3)
    r += 1
    for row in rows:
        if isinstance(row, Group):
            cell = ws.cell(r, 1, ("    " * row.level) + row.label)
            cell.font = Font(bold=True, size=12, color=("FFFFFF" if row.level == 0 else "3B0764"))
            for c in range(1, ncols + 1):
                ws.cell(r, c).fill = _XL_DEPT if row.level == 0 else _XL_DESIG
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        else:
            max_lines = 1
            for c, v in enumerate(row, 1):
                val = _xl_value(v)
                cell = ws.cell(r, c, val)
                cell.font = _XL_BODY
                multiline = isinstance(val, str) and "\n" in val
                if multiline:
                    max_lines = max(max_lines, val.count("\n") + 1)
                if c > 2 or multiline:
                    cell.alignment = Alignment(horizontal="center", vertical="center",
                                               wrap_text=multiline)
            if max_lines > 1:  # stacked IN/OUT/hours cells need room
                ws.row_dimensions[r].height = 13.5 * max_lines + 2
        r += 1
    widths = [max(len(str(h)), 10) for h in headers]
    for row in rows:
        if isinstance(row, Group):
            continue
        for i, v in enumerate(row[:ncols]):
            lv = max((len(x) for x in str(v or "").split("\n")), default=0)
            widths[i] = min(max(widths[i], lv), 40)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = int(w * 1.1) + 2  # wider text needs wider columns
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF

_PURPLE = colors.HexColor("#5b21b6")
_LIGHT = colors.HexColor("#ede9fe")
_DEEP = colors.HexColor("#4c1d95")
# Style for multi-line table cells (stacked IN / OUT).
def _cell_style(delta: float = 0) -> ParagraphStyle:
    return ParagraphStyle("cell", fontName="Helvetica", fontSize=6 + delta, leading=6.8 + delta, alignment=TA_CENTER)


def _pdf(title: str, subtitle: str, headers: list[str], rows: list[Any],
         col_widths: list[float] | None = None, org: dict | None = None) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            topMargin=12 * mm, bottomMargin=10 * mm,
                            leftMargin=10 * mm, rightMargin=10 * mm, title=title)
    styles = getSampleStyleSheet()
    story = []
    if org:
        org_name = org.get("OrgName") or ""
        org_addr = " · ".join(p for p in [org.get("OrgAddress"), org.get("OrgPhone"), org.get("OrgEmail")] if p)
        if org_name:
            story.append(Paragraph(org_name, ParagraphStyle("org", parent=styles["Normal"],
                                   fontSize=13, textColor=_PURPLE, fontName="Helvetica-Bold")))
        if org_addr:
            story.append(Paragraph(org_addr, ParagraphStyle("orgA", parent=styles["Normal"],
                                   fontSize=9, textColor=colors.grey)))
        if org_name or org_addr:
            story.append(Spacer(1, 4))
    story += [
        Paragraph(title, ParagraphStyle("t", parent=styles["Title"], textColor=_PURPLE, fontSize=16)),
        Paragraph(subtitle, ParagraphStyle("s", parent=styles["Normal"], fontSize=10, textColor=colors.grey)),
        Spacer(1, 6),
    ]
    ncols = len(headers)

    def build_table(delta: float) -> Table:
        """Table text is `delta` points larger than the original sizes."""
        cell_style = _cell_style(delta)
        table_data: list[list[Any]] = [headers]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), _DEEP),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 6.5 + delta),
            ("FONT", (0, 1), (-1, -1), "Helvetica", 6.5 + delta),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
            ("ALIGN", (2, 1), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ]
        stripe = 0
        for r in rows:
            i = len(table_data)
            if isinstance(r, Group):
                table_data.append([("    " * r.level) + r.label] + [""] * (ncols - 1))
                style += [("SPAN", (0, i), (-1, i)), ("ALIGN", (0, i), (-1, i), "LEFT"),
                          ("FONT", (0, i), (-1, i), "Helvetica-Bold", (7.5 if r.level == 0 else 7) + delta),
                          ("BACKGROUND", (0, i), (-1, i), _PURPLE if r.level == 0 else _LIGHT),
                          ("TEXTCOLOR", (0, i), (-1, i), colors.white if r.level == 0 else _DEEP)]
                stripe = 0
            else:
                cells = []
                for c in r:
                    s = "" if c is None else str(c)
                    cells.append(Paragraph(s.replace("\n", "<br/>"), cell_style) if "\n" in s else s)
                table_data.append(cells)
                if stripe % 2:
                    style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#f7f5ff")))
                stripe += 1
        tbl = Table(table_data, repeatRows=1, colWidths=col_widths)
        tbl.setStyle(TableStyle(style))
        return tbl

    # One point larger than before wherever the table still fits across the page (the wide monthly grids are
    # already full width); otherwise fall back a half point at a time, to the original size at worst.
    tbl = build_table(0)
    for delta in (1.0, 0.5):
        candidate = build_table(delta)
        if candidate.wrap(doc.width, 10_000)[0] <= doc.width:
            tbl = candidate
            break
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- Report builders
# Each returns (title, subtitle, headers, rows); the render_* functions below pick the format.

def _register(data: dict, grouped: bool):
    days = data["DaysInMonth"]
    leave_codes = data["LeaveCodes"]
    headers = (["S.No", "Employee Code", "Employee Name", "Department", "Designation"] + [str(d) for d in range(1, days + 1)]
               + ["Present", "Absent", "Half", "WeeklyOff", "Holiday"] + leave_codes)
    rows = _rows(data, lambda e: (
        [e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Department"] or "", e.get("Designation") or ""]
        + e["days"]
        + [e["PresentDays"], e["AbsentDays"], e["HalfDays"], e["WeeklyOffDays"], e["HolidayDays"]]
        + [e["LeaveByCode"].get(code, 0) for code in leave_codes]), grouped)
    return "Monthly Attendance Register", data["MonthLabel"], headers, rows


def _leave_balance(data: dict, grouped: bool):
    types = data["LeaveTypes"]
    headers = ["S.No", "Employee Code", "Employee Name", "Department", "Designation"]
    for t in types:
        headers += [f"{t['Code']} Entitled", f"{t['Code']} Taken", f"{t['Code']} Balance"]

    def row(e: dict) -> list[Any]:
        by_code = {b["Code"]: b for b in e["balances"]}
        out = [e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Department"] or "", e.get("Designation") or ""]
        for t in types:
            b = by_code.get(t["Code"], {"Entitled": 0, "Taken": 0, "Balance": 0})
            out += [b["Entitled"], b["Taken"], b["Balance"]]
        return out
    return "Leave Balance Report", f"Year {data['Year']}", headers, _rows(data, row, grouped)


def _daily(data: dict, grouped: bool):
    kind = {"present": "Present List", "absent": "Absent List", "leave": "On Leave"}.get(data["Kind"], "Daily Attendance")
    t = data["Totals"]
    subtitle = (f"{data['DayLabel']}  —  Present {t['present']} (Late {t.get('late', 0)}), "
                f"Absent {t['absent']}, Leave {t['leave']}, Off {t['off']} of {t['total']}")
    headers = ["S.No", "Employee Code", "Employee Name", "Department", "Designation", "Status", "Category", "In", "Out", "Late (min)", "Remarks"]
    rows = _rows(data, lambda e: [
        e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Department"] or "", e.get("Designation") or "",
        (e["LeaveCode"] or e["StatusLabel"]) if e["Status"] == "L" else e["StatusLabel"],
        e["CategoryName"] or "", e["FirstIn"] or "", e["LastOut"] or "", e["LateMinutes"] or 0, e["Remarks"] or ""], grouped)
    return kind, subtitle, headers, rows


def _late_ot(data: dict, grouped: bool = False):
    # Overtime is not reported (removed at the college's request, Oct 2026); the data still carries it.
    headers = ["S.No", "Employee Code", "Employee Name", "Department", "Late (min)", "Early (min)", "Late Days"]
    rows = [[e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Department"] or "",
             e["TotalLateMinutes"], e["TotalEarlyMinutes"], e["LateDays"]]
            for e in data["employees"]]
    return "Late / Early Report", data["MonthLabel"], headers, rows


def _individual(data: dict, grouped: bool = False):
    headers = ["Date", "Day", "Status", "Leave", "Category", "In", "Out", "Late", "Early", "Remarks"]
    rows = [[d["Date"], d["DayLabel"], d["StatusLabel"], d["LeaveCode"] or "", d["CategoryName"] or "",
             d["FirstIn"] or "", d["LastOut"] or "", d["Late"], d["Early"],
             d["Remarks"] or ""] for d in data["days"]]
    subtitle = f"{data['EmployeeName'] or data['EmployeeCode']} — {data['MonthLabel']}"
    return "Individual Attendance Sheet", subtitle, headers, rows


def _range_subtitle(data: dict) -> str:
    return f"From {data['FromLabel']} To {data['ToLabel']}    Category : {data['Category']}"


def _muster_time(data: dict, grouped: bool):
    headers = ["S.No", "Employee Code", "Employee Name"] + data["DateHeaders"] + ["Late Hrs"]
    rows = _rows(data, lambda e: (
        [e["SNo"], e["EmployeeCode"], e["EmployeeName"]] + e["days"] + [e["LateHrs"]]), grouped)
    return "Attendance Report", _range_subtitle(data), headers, rows


def _muster_leave(data: dict, grouped: bool):
    cols = data["Columns"]
    headers = (["S.No", "Employee Code", "Employee Name"] + data["DateHeaders"] + cols
               + ["Total", "Late Hrs", "Late Cl"])
    rows = _rows(data, lambda e: (
        [e["SNo"], e["EmployeeCode"], e["EmployeeName"]] + e["days"]
        + [e["Counts"].get(c, 0) for c in cols]
        + [e["Total"], e["LateHrs"], e["LateCl"]]), grouped)
    return "Attendance Report (Time & Leave)", _range_subtitle(data), headers, rows


def _late_arrival(data: dict, grouped: bool):
    headers = ["S.No", "Employee Code", "Employee Name", "Date", "Shift", "Actual In", "Late Hrs"]
    rows = _rows(data, lambda e: [
        e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Date"],
        e["ShiftTime"], e["FirstIn"], e["Minutes"]], grouped)
    return "Late Arrival Report", _range_subtitle(data), headers, rows


def _early_departure(data: dict, grouped: bool):
    headers = ["S.No", "Employee Code", "Employee Name", "Date", "Shift", "Actual Out", "Early Hrs"]
    rows = _rows(data, lambda e: [
        e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Date"],
        e["ShiftTime"], e["LastOut"], e["Minutes"]], grouped)
    return "Early Departure Report", _range_subtitle(data), headers, rows


def _in_out(data: dict, grouped: bool):
    headers = ["S.No", "Employee Code", "Employee Name", "Date", "Punches", "Count"]
    rows = _rows(data, lambda e: [
        e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Date"], e["Punches"], e["Count"]], grouped)
    return "In-Out List", _range_subtitle(data), headers, rows


def _employee_list(data: dict, grouped: bool):
    headers = ["S.No", "Employee Code", "Employee Name", "Department", "Designation",
               "Type", "Category", "Status", "Joining Date"]
    rows = _rows(data, lambda e: [
        e["SNo"], e["EmployeeCode"], e["EmployeeName"], e["Department"] or "", e.get("Designation") or "",
        e["Type"], e["CategoryName"], e["Status"], e["JoiningDate"]], grouped)
    subtitle = f"{data['Count']} employees    Category : {data['Category']}    as on {data['GeneratedLabel']}"
    return "Employee List", subtitle, headers, rows


_BUILDERS = {
    "register": _register, "leave-balance": _leave_balance, "daily": _daily,
    "late-ot": _late_ot, "individual": _individual,
    "muster-time": _muster_time, "muster-leave": _muster_leave,
    "late-arrival": _late_arrival, "early-departure": _early_departure,
    "in-out": _in_out, "employee-list": _employee_list,
}

MEDIA_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def render(report: str, fmt: str, data: dict, grouped: bool = False, org: dict | None = None) -> bytes | str:
    """Render `report` ('register', 'leave-balance', 'daily', 'late-ot', 'individual')
    as csv (str), xlsx or pdf (bytes)."""
    if report in ("muster-time", "muster-leave") and fmt in ("pdf", "xlsx"):   # laid out like the college's printed form
        from app.services import muster_format
        make = muster_format.muster_pdf if fmt == "pdf" else muster_format.muster_xlsx
        return make(data, org, with_leave=(report == "muster-leave"), grouped=grouped)
    title, subtitle, headers, rows = _BUILDERS[report](data, grouped)
    if data.get("AttendanceCategory"):   # the attendance-category filter, shown on every report's header
        subtitle += f"    Attendance Category : {data['AttendanceCategory']}"
    if fmt == "pdf":
        return _pdf(title, subtitle, headers, rows, org=org)
    if fmt == "xlsx":
        return _xlsx(title, subtitle, headers, rows, org=org)
    return _csv(headers, rows)


def simple_table(title: str, subtitle: str, headers: list[str], rows: list[list[Any]],
                 fmt: str, org: dict | None = None) -> bytes | str:
    """Render an arbitrary flat table (each row a list of cells) with the shared report
    styling (purple header, striped rows, org stamp) used by the employee reports."""
    if fmt == "pdf":
        return _pdf(title, subtitle, headers, rows, org=org)
    if fmt == "xlsx":
        return _xlsx(title, subtitle, headers, rows, org=org)
    return _csv(headers, rows)


# Backwards-compatible names used by older callers.
def register_csv(data): return render("register", "csv", data)
def leave_balance_csv(data): return render("leave-balance", "csv", data)
def late_ot_csv(data): return render("late-ot", "csv", data)
def individual_csv(data): return render("individual", "csv", data)
def register_pdf(data, org=None): return render("register", "pdf", data, org=org)
def leave_balance_pdf(data, org=None): return render("leave-balance", "pdf", data, org=org)
def late_ot_pdf(data, org=None): return render("late-ot", "pdf", data, org=org)
def individual_pdf(data, org=None): return render("individual", "pdf", data, org=org)
