"""Bulk employee import from an Excel workbook (Application tier).

Two entry points used by the API:
- inspect(): parse the workbook just enough to power the column-mapping UI
  (detect the header row, list columns with sample values, guess a mapping).
- process(): given a confirmed mapping, validate every data row and — when
  commit=True — upsert employees (by EmployeeCode) and auto-create any missing
  Organization / Department / Designation / EmployeeType masters.

Nothing is written unless commit=True; process(commit=False) is the dry-run
that backs the preview screen.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from sqlalchemy.orm import Session

from app.models.entities import (
    DepartmentMaster,
    Designation,
    Employee,
    EmployeeType,
    Organization,
    School,
)
from app.schemas.employee import (
    IMPORT_FIELDS,
    ImportColumn,
    ImportInspectResult,
    ImportResult,
    ImportRowResult,
    ImportSummary,
)

# How many rows from the top to search for the header, and how many sample
# values to collect per column for the mapping UI.
MAX_HEADER_SCAN_ROWS = 25
SAMPLE_COUNT = 3

# field -> set of normalized header aliases used to guess the mapping.
_ALIASES: dict[str, set[str]] = {
    "EmployeeCode": {"employeeno", "employeecode", "empno", "empcode", "code", "employeeid", "userid", "faceid", "deviceid", "biometricid"},
    "HrCode": {"hrcode", "hrno", "hremployeecode", "employeenumber", "staffno", "staffcode"},
    # Shown as "Sort No" (the office's ordering number); older sheets say "Short No".
    "ShortNo": {"sortno", "sortnumber", "sortorder", "shortno", "shortnumber", "shortcode", "shortid", "shortnum"},
    "EmployeeName": {"employeename", "name", "empname", "fullname"},
    "EmployeeRFIDNumber": {"accesscardno", "cardno", "rfid", "rfidno", "accesscard", "cardnumber"},
    "DateOfBirth": {"dob", "dateofbirth", "birthdate"},
    "JoiningDate": {"doj", "dateofjoining", "joiningdate", "joindate"},
    "Organization": {"company", "organization", "organisation", "institution", "org"},
    "Department": {"department", "dept"},
    "Designation": {"designation", "role", "post", "position"},
    "EmployeeType": {"employeetype", "type", "category", "empanelment"},
    "School": {"school", "schoolname", "faculty"},
    "CurrentStatus": {"status", "currentstatus", "activestatus"},
}

_FACE_HEADERS = {"faceid", "deviceid", "biometricid"}
_HR_NUMBER_HEADERS = {"employeeno", "employeecode", "empno", "empcode", "employeenumber"}

_STATUS_MAP = {
    "active": "Active",
    "inactive": "Inactive",
    "resigned": "Resigned",
    "terminated": "Resigned",  # "Terminated" is no longer offered; treat it as an exit
}

_DATE_FORMATS = ("%d-%b-%Y", "%d-%b-%y", "%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y")


def _norm(text: Any) -> str:
    """Normalize a header cell for alias matching: lower-case, drop non-alnum."""
    return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _cell_str(value: Any) -> str:
    """Render a cell value as a trimmed display string."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _to_iso_date(value: Any) -> Optional[str]:
    """Return YYYY-MM-DD, or None if the value is empty / unparseable."""
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    if not text:
        return None
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None  # signal "could not parse" to the caller


class EmployeeImportService:
    def __init__(self, session: Session) -> None:
        self._session = session

    # ---- workbook loading ----

    def _rows(self, file_bytes: bytes) -> list[list[Any]]:
        wb = load_workbook(BytesIO(file_bytes), data_only=True, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows

    def _detect_header_row(self, rows: list[list[Any]]) -> int:
        """Return the 1-based index of the row that best looks like a header."""
        all_aliases = set().union(*_ALIASES.values())
        best_idx, best_score = 0, -1
        for i, row in enumerate(rows[:MAX_HEADER_SCAN_ROWS]):
            score = sum(1 for c in row if _norm(c) in all_aliases)
            if score > best_score:
                best_idx, best_score = i, score
        return best_idx + 1  # 1-based

    def _guess_mapping(self, header_cells: list[Any]) -> dict[str, int]:
        mapping: dict[str, int] = {}
        for idx, cell in enumerate(header_cells):
            key = _norm(cell)
            if not key:
                continue
            for field, aliases in _ALIASES.items():
                if field not in mapping and key in aliases:
                    mapping[field] = idx
        # A sheet with a Face ID column AND an "Employee Code / No" column: the first is the
        # device key, the second is the HR code shown in lists. (Without a Face ID column the
        # "Employee Code" column keeps being the key, as before.)
        face_idx = mapping.get("EmployeeCode")
        if face_idx is not None and "HrCode" not in mapping and _norm(header_cells[face_idx]) in _FACE_HEADERS:
            for idx, cell in enumerate(header_cells):
                if idx != face_idx and _norm(cell) in _HR_NUMBER_HEADERS:
                    mapping["HrCode"] = idx
                    break
        return mapping

    # ---- public: inspect for the mapping UI ----

    def inspect(self, file_bytes: bytes) -> ImportInspectResult:
        rows = self._rows(file_bytes)
        if not rows:
            return ImportInspectResult(header_row=1, columns=[], guessed_mapping={})
        header_row = self._detect_header_row(rows)
        header_cells = rows[header_row - 1]
        ncols = max((len(r) for r in rows), default=len(header_cells))

        columns: list[ImportColumn] = []
        for idx in range(ncols):
            header = _cell_str(header_cells[idx]) if idx < len(header_cells) else ""
            samples: list[str] = []
            for r in rows[header_row:]:
                if idx < len(r):
                    s = _cell_str(r[idx])
                    if s:
                        samples.append(s)
                if len(samples) >= SAMPLE_COUNT:
                    break
            columns.append(
                ImportColumn(index=idx, letter=get_column_letter(idx + 1), header=header or None, samples=samples)
            )
        return ImportInspectResult(
            header_row=header_row,
            columns=columns,
            guessed_mapping=self._guess_mapping(header_cells),
            fields=IMPORT_FIELDS,
        )

    # ---- master get-or-create (only called when committing) ----

    def _get_or_create_org(self, name: str, created: list[str]) -> Optional[int]:
        name = (name or "").strip()
        if not name:
            return None
        row = self._session.query(Organization).filter(Organization.OrgName == name).first()
        if not row:
            row = Organization(OrgName=name, IsActive="Y")
            self._session.add(row)
            self._session.flush()
            created.append(name)
        return row.OrganizationId

    def _default_org_id(self) -> int:
        row = self._session.query(Organization).order_by(Organization.OrganizationId).first()
        if row:
            return row.OrganizationId
        row = Organization(OrgName="Default Organization", IsActive="Y")
        self._session.add(row)
        self._session.flush()
        return row.OrganizationId

    def _ensure_department(self, org_id: int, name: str, created: list[str]) -> None:
        name = (name or "").strip()
        if not name:
            return
        exists = (
            self._session.query(DepartmentMaster)
            .filter(DepartmentMaster.OrganizationId == org_id)
            .filter(DepartmentMaster.DepartmentName == name)
            .first()
        )
        if not exists:
            self._session.add(DepartmentMaster(OrganizationId=org_id, DepartmentName=name, IsActive="Y"))
            self._session.flush()
            created.append(name)

    def _ensure_designation(self, name: str, created: list[str]) -> None:
        name = (name or "").strip()
        if not name:
            return
        if not self._session.query(Designation).filter(Designation.DesignationName == name).first():
            self._session.add(Designation(DesignationName=name, IsActive="Y"))
            self._session.flush()
            created.append(name)

    def _get_or_create_type(self, name: str, created: list[str]) -> Optional[int]:
        name = (name or "").strip()
        if not name:
            return None
        row = self._session.query(EmployeeType).filter(EmployeeType.TypeName == name).first()
        if not row:
            row = EmployeeType(TypeName=name, IsActive="Y")
            self._session.add(row)
            self._session.flush()
            created.append(name)
        return row.EmployeeTypeId

    def _get_or_create_school(self, name: str, created: list[str]) -> Optional[int]:
        name = (name or "").strip()
        if not name:
            return None
        row = self._session.query(School).filter(School.SchoolName == name).first()
        if not row:
            row = School(SchoolName=name, IsActive="Y")
            self._session.add(row)
            self._session.flush()
            created.append(name)
        return row.SchoolId

    # ---- public: validate / commit ----

    def process(self, file_bytes: bytes, mapping: dict[str, int], header_row: int, commit: bool) -> ImportResult:
        rows = self._rows(file_bytes)
        result = ImportResult(committed=False, summary=ImportSummary(), rows=[])
        masters = {"organizations": [], "departments": [], "designations": [], "employee_types": [], "schools": []}

        if "EmployeeCode" not in mapping or "EmployeeName" not in mapping:
            # Nothing sensible to do without at least code + name mapped.
            return result

        # Existing codes (for upsert action) and the running NumericCode counter.
        existing_codes = {e.EmployeeCode for e in self._session.query(Employee.EmployeeCode).all()}
        max_numeric = self._session.query(Employee.NumericCode).order_by(Employee.NumericCode.desc()).first()
        next_numeric = (max_numeric[0] if max_numeric and max_numeric[0] else 0)
        renamed: list[Employee] = []   # existing people whose name this import changes

        # RFID must be unique across employees ('' and '0' mean "no card"). Seed the
        # owner map from the DB, then claim numbers row-by-row so duplicates inside the
        # file are caught too. A conflicting card is left blank with a warning.
        rfid_owner: dict[str, str] = {}
        for r, c in self._session.query(Employee.EmployeeRFIDNumber, Employee.EmployeeCode).all():
            v = (r or "").strip()
            if v and v != "0":
                rfid_owner.setdefault(v, c)

        # HR codes are unique too (shown in lists/reports); claim them row-by-row like RFID cards.
        hr_owner: dict[str, str] = {}
        for h, c in self._session.query(Employee.StringCode, Employee.EmployeeCode).all():
            if h:
                hr_owner.setdefault(h.strip(), c)

        def cell(row: list[Any], field: str) -> Any:
            idx = mapping.get(field)
            if idx is None or idx >= len(row):
                return None
            return row[idx]

        data_rows = rows[header_row:]  # everything below the header
        for offset, row in enumerate(data_rows):
            row_number = header_row + offset + 1  # 1-based spreadsheet row
            if not any(_cell_str(v) for v in row):
                continue  # fully blank

            code = _cell_str(cell(row, "EmployeeCode"))
            name = _cell_str(cell(row, "EmployeeName"))
            dept = _cell_str(cell(row, "Department"))
            desig = _cell_str(cell(row, "Designation"))
            emp_type = _cell_str(cell(row, "EmployeeType"))
            school_name = _cell_str(cell(row, "School"))
            org_name = _cell_str(cell(row, "Organization"))
            rfid = _cell_str(cell(row, "EmployeeRFIDNumber"))
            hr_code = _cell_str(cell(row, "HrCode"))
            short_no = _cell_str(cell(row, "ShortNo"))[:20]
            messages: list[str] = []

            if not code or not name:
                result.rows.append(
                    ImportRowResult(row_number=row_number, EmployeeCode=code or None, EmployeeName=name or None,
                                    action="error", messages=["Missing Employee Code or Name."]))
                result.summary.errors += 1
                result.summary.total += 1
                continue

            # Dates (blank if unparseable, with a warning).
            dob = None
            if "DateOfBirth" in mapping:
                raw = cell(row, "DateOfBirth")
                if _cell_str(raw):
                    dob = _to_iso_date(raw)
                    if dob is None:
                        messages.append(f"Could not read Date of Birth '{_cell_str(raw)}'.")
            doj = None
            if "JoiningDate" in mapping:
                raw = cell(row, "JoiningDate")
                if _cell_str(raw):
                    doj = _to_iso_date(raw)
                    if doj is None:
                        messages.append(f"Could not read Joining Date '{_cell_str(raw)}'.")

            status = _STATUS_MAP.get(_cell_str(cell(row, "CurrentStatus")).lower(), None) if "CurrentStatus" in mapping else None

            # RFID uniqueness: claim the card for this code, or drop it with a warning.
            rfid_norm = rfid.strip()
            rfid_ok = False
            if rfid_norm and rfid_norm != "0":
                owner = rfid_owner.get(rfid_norm)
                if owner and owner != code:
                    messages.append(f"RFID card {rfid_norm} already used by {owner}; left blank.")
                else:
                    rfid_owner[rfid_norm] = code
                    rfid_ok = True

            # HR code: keep it only if nobody else already has it, else fall back to the Face ID.
            hr_ok = False
            if hr_code:
                owner = hr_owner.get(hr_code)
                if owner and owner != code:
                    messages.append(f"Employee code {hr_code} already used by {owner}; kept the Face ID instead.")
                else:
                    hr_owner[hr_code] = code
                    hr_ok = True

            action = "update" if code in existing_codes else "create"
            result.summary.total += 1
            if action == "create":
                result.summary.to_create += 1
            else:
                result.summary.to_update += 1

            result.rows.append(
                ImportRowResult(
                    row_number=row_number, EmployeeCode=code, EmployeeName=name,
                    Department=dept or None, Designation=desig or None,
                    EmployeeType=emp_type or None, CurrentStatus=status,
                    action=action, messages=messages,
                )
            )

            if not commit:
                continue

            # --- write path ---
            org_id = self._get_or_create_org(org_name, masters["organizations"]) if org_name else None
            if dept:
                self._ensure_department(org_id or self._default_org_id(), dept, masters["departments"])
            if desig:
                self._ensure_designation(desig, masters["designations"])
            type_id = self._get_or_create_type(emp_type, masters["employee_types"]) if emp_type else None
            school_id = self._get_or_create_school(school_name, masters["schools"]) if school_name else None

            emp = self._session.query(Employee).filter(Employee.EmployeeCode == code).first()
            if not emp:
                # NumericCode: card if positive int, else code if all-digits, else next sequential.
                if rfid.isdigit() and int(rfid) > 0:
                    numeric = int(rfid)
                elif code.isdigit():
                    numeric = int(code)
                else:
                    next_numeric += 1
                    numeric = next_numeric
                emp = Employee(
                    EmployeeCode=code, StringCode=hr_code if hr_ok else code, NumericCode=numeric,
                    EmployeeName=name, CurrentStatus=status or "Active",
                )
                from app.services.employee_delete import unblock_pin
                unblock_pin(self._session, code)   # importing the person again lifts a block from an earlier delete
                self._session.add(emp)
                existing_codes.add(code)
            else:
                if commit and name and emp.EmployeeName != name:
                    renamed.append(emp)   # sent to the readers this person is on, after the save
                emp.EmployeeName = name
                if status:
                    emp.CurrentStatus = status

            # Set optional fields only when provided (don't blank existing data).
            if hr_ok:
                emp.StringCode = hr_code
            if short_no:   # blank cell keeps what is already saved
                emp.ShortNo = short_no
            if rfid_ok:
                emp.EmployeeRFIDNumber = rfid_norm
            if dept:
                emp.Department = dept
            if desig:
                emp.Designation = desig
            if dob:
                emp.DateOfBirth = dob
            if doj:
                emp.JoiningDate = doj
            if org_id is not None:
                emp.OrganizationId = org_id
            if type_id is not None:
                emp.EmployeeTypeId = type_id
            if school_id is not None:
                emp.SchoolId = school_id

        if commit:
            try:
                self._session.commit()
                result.committed = True
                result.summary.created = result.summary.to_create
                result.summary.updated = result.summary.to_update
            except Exception as exc:  # pragma: no cover - surfaced to the UI
                self._session.rollback()
                raise exc
            if renamed:
                # New names go to the readers these people are on; the import itself is already saved.
                from app.services.push_channel import queue_name_change
                try:
                    for emp in renamed:
                        queue_name_change(self._session, emp)
                    self._session.commit()
                except Exception:  # noqa: BLE001 - readers can still be updated with Send names
                    self._session.rollback()

        result.masters_created = masters
        return result
